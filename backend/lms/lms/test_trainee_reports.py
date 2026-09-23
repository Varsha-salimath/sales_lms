import base64
import io

import frappe
import openpyxl
from frappe.tests import UnitTestCase

from lms.lms.trainee_reports import (
	export_active_trainees_csv,
	export_exit_churn_report,
	export_finance_reconciliation_report,
	export_location_cohort_report,
	export_training_outcome_report,
	export_vendor_payroll_input,
	export_weekly_attendance_payroll_report,
)


class TestTraineeReports(UnitTestCase):
	def _open_workbook(self, encoded):
		return openpyxl.load_workbook(io.BytesIO(base64.b64decode(encoded)))

	def test_export_includes_only_in_training_trainees(self):
		active = frappe.get_doc(
			{
				"doctype": "Sales Trainee",
				"trainee_name": "Active One",
				"personal_email": f"active-{frappe.generate_hash(length=6)}@example.com",
				"date_of_joining": "2026-10-06",
				"training_status": "In Training",
			}
		).insert(ignore_permissions=True)
		exited = frappe.get_doc(
			{
				"doctype": "Sales Trainee",
				"trainee_name": "Exited One",
				"personal_email": f"exited-{frappe.generate_hash(length=6)}@example.com",
				"date_of_joining": "2026-10-06",
				"training_status": "Resigned",
			}
		).insert(ignore_permissions=True)

		encoded = export_active_trainees_csv()
		workbook = openpyxl.load_workbook(io.BytesIO(base64.b64decode(encoded)))
		sheet = workbook.active
		names_in_sheet = {row[0].value for row in sheet.iter_rows(min_row=2)}
		self.assertIn(active.trainee_name, names_in_sheet)
		self.assertNotIn(exited.trainee_name, names_in_sheet)

	def _get_or_create_user(self, email):
		if frappe.db.exists("User", email):
			return email
		frappe.get_doc({"doctype": "User", "email": email, "first_name": email.split("@")[0]}).insert(
			ignore_permissions=True
		)
		return email

	def test_user_without_required_role_is_rejected(self):
		# Security Review Focus: export_active_trainees_csv uses frappe.get_all,
		# which bypasses permission_query_conditions/has_permission hooks. A bare
		# @frappe.whitelist() with no role gate would let any authenticated user
		# dump PII (personal emails, locations) for all active trainees.
		email = self._get_or_create_user(f"report-no-access-{frappe.generate_hash(length=6)}@example.com")
		frappe.set_user(email)
		self.addCleanup(frappe.set_user, "Administrator")
		with self.assertRaises(frappe.PermissionError):
			export_active_trainees_csv()

	def test_new_report_export_rejects_user_without_required_role(self):
		# Representative coverage of the shared _ensure_report_access() gate on one of
		# the 6 new exports; the helper itself is already covered above for
		# export_active_trainees_csv.
		email = self._get_or_create_user(f"report-no-access-{frappe.generate_hash(length=6)}@example.com")
		frappe.set_user(email)
		self.addCleanup(frappe.set_user, "Administrator")
		with self.assertRaises(frappe.PermissionError):
			export_exit_churn_report()

	def test_export_exit_churn_report_includes_only_exited(self):
		frappe.get_doc(
			{
				"doctype": "Sales Trainee",
				"trainee_name": "Churned",
				"personal_email": f"churn-{frappe.generate_hash(length=6)}@example.com",
				"date_of_joining": "2026-10-06",
				"training_status": "Absconded",
				"exit_date": "2026-10-09",
				"exit_reason": "Did not return after weekend",
			}
		).insert(ignore_permissions=True)
		frappe.get_doc(
			{
				"doctype": "Sales Trainee",
				"trainee_name": "Still Training",
				"personal_email": f"active-churn-{frappe.generate_hash(length=6)}@example.com",
				"date_of_joining": "2026-10-06",
				"training_status": "In Training",
			}
		).insert(ignore_permissions=True)
		workbook = self._open_workbook(export_exit_churn_report())
		names = {row[0].value for row in workbook.active.iter_rows(min_row=2)}
		self.assertIn("Churned", names)
		self.assertNotIn("Still Training", names)

	def test_export_training_outcome_report_groups_by_status(self):
		frappe.get_doc(
			{
				"doctype": "Sales Trainee",
				"trainee_name": "Cleared One",
				"personal_email": f"outcome-{frappe.generate_hash(length=6)}@example.com",
				"date_of_joining": "2026-10-06",
				"training_status": "Training Cleared",
			}
		).insert(ignore_permissions=True)
		workbook = self._open_workbook(export_training_outcome_report())
		statuses = {row[1].value for row in workbook.active.iter_rows(min_row=2)}
		self.assertIn("Training Cleared", statuses)

	def test_export_vendor_payroll_input_lists_cycle_rows(self):
		trainee = frappe.get_doc(
			{
				"doctype": "Sales Trainee",
				"trainee_name": "Vendor Row",
				"personal_email": f"vendor-{frappe.generate_hash(length=6)}@example.com",
				"date_of_joining": "2026-10-06",
				"salary": 20000,
			}
		).insert(ignore_permissions=True)
		cycle = frappe.get_doc(
			{
				"doctype": "Weekly Payroll Cycle",
				"week_start": "2026-10-06",
				"week_end": "2026-10-12",
				"payroll_inputs": [{"trainee": trainee.name, "salary": 20000, "working_days": 6, "payable_days": 5}],
			}
		).insert(ignore_permissions=True)
		workbook = self._open_workbook(export_vendor_payroll_input(cycle.name))
		rows = list(workbook.active.iter_rows(min_row=2, values_only=True))
		self.assertEqual(len(rows), 1)
		self.assertEqual(rows[0][0], trainee.name)

	def test_export_finance_reconciliation_report_shows_captured_vs_invoiced(self):
		cycle = frappe.get_doc(
			{
				"doctype": "Weekly Payroll Cycle",
				"week_start": "2026-10-06",
				"week_end": "2026-10-12",
				"vendor_invoice_amount": 95000,
			}
		).insert(ignore_permissions=True)
		workbook = self._open_workbook(export_finance_reconciliation_report(cycle.name))
		row = next(workbook.active.iter_rows(min_row=2, values_only=True))
		self.assertEqual(row[-1], 95000)

	def test_export_location_cohort_report_lists_every_cohort(self):
		frappe.get_doc(
			{"doctype": "Sales Trainee Cohort", "cohort_name": f"Loc-{frappe.generate_hash(length=6)}", "location": "Chennai", "start_date": "2026-10-06"}
		).insert(ignore_permissions=True)
		workbook = self._open_workbook(export_location_cohort_report())
		locations = {row[1].value for row in workbook.active.iter_rows(min_row=2)}
		self.assertIn("Chennai", locations)
